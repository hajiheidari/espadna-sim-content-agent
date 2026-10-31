import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
import openpyxl
from espadna.core import bootstrap_owner, connect, set_role, set_sale_status, caption, phone_text, phone_values
from espadna.importer import (apply_saved_sync, import_workbook, parse_workbook,
                              prepare_saved_sync, save_sync_upload)

# Synthetic format-only identifiers; no production workbook or customer records.
PHONE = '09000000000'
OTHER = '09000000001'
HEADER = [' شماره دلخواه', 'نوع سیم کارت', 'وضعیت سیم', 'قیمت', 'شرایط فروش ', 'اپراتور']


class InventoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'test.xlsx'
        self.db = connect(':memory:')
        bootstrap_owner(self.db, 1)

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def book(self, rows=None, second=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADER)
        for row in rows if rows is not None else [[PHONE, 1, 2, 12345, '=ignored()', 'همراه اول']]:
            ws.append(row)
        if second is not None:
            ws = wb.create_sheet('second')
            ws.append(HEADER)
            for row in second:
                ws.append(row)
        wb.save(self.path)
        wb.close()
        return self.path

    def test_incremental_idempotent_no_delete_no_sale_reset(self):
        self.book([[PHONE, 1, 2, 12345, 1, 'همراه اول'], [OTHER, 2, 1, 0, 2, None]])
        self.assertEqual(import_workbook(self.db, 1, self.path)['inserted'], 2)
        set_sale_status(self.db, 1, PHONE, 'sold')
        self.book()
        self.assertEqual(import_workbook(self.db, 1, self.path)['unchanged'], 1)
        self.book([[PHONE, 2, 3, 999, 'anything', None]])
        self.assertEqual(import_workbook(self.db, 1, self.path)['updated'], 1)
        row = self.db.execute('SELECT * FROM inventory WHERE phone_normalized=?', (PHONE,)).fetchone()
        self.assertEqual((row['sale_status'], row['sim_status'], row['phone_normalized']), ('sold', 3, PHONE))
        self.assertEqual(self.db.execute('SELECT count(*) FROM inventory').fetchone()[0], 2)

    def test_duplicate_across_sheets(self):
        row = [PHONE, 1, 2, 12345, 1, 'همراه اول']
        self.book([row], [row])
        self.assertEqual(import_workbook(self.db, 1, self.path)['duplicates'], 1)
        row2 = list(row)
        row2[1] = 2
        self.book([row], [row2])
        with self.assertRaises(ValueError):
            import_workbook(self.db, 1, self.path)
        self.assertEqual(self.db.execute('SELECT sim_type FROM inventory').fetchone()[0], 1)

    def test_full_sync_preview_apply_backup_and_sale_status_guards(self):
        db_path = Path(self.tmp.name) / 'full.db'
        db = connect(db_path)
        bootstrap_owner(db, 1)
        set_role(db, 1, 2, 'admin')
        try:
            third = '09000000002'
            new = '09000000003'
            self.book([[PHONE, 1, 2, 10, None, 'همراه اول'],
                       [OTHER, 1, 2, 20, None, 'همراه اول'],
                       [third, 2, 1, 30, None, 'ایرانسل']])
            import_workbook(db, 1, self.path)
            set_sale_status(db, 1, PHONE, 'sold')
            set_sale_status(db, 1, OTHER, 'reserved')
            self.book([['0900 000 0000', 2, 3, 11, None, 'همراه اول'],
                       [OTHER, 1, 2, 20, None, 'همراه اول'],
                       [new, 2, 1, 40, None, 'رایتل']])
            sync_id = save_sync_upload(db, 1, self.path)
            report = prepare_saved_sync(db, 1, sync_id, 'full')
            self.assertEqual({k: report[k] for k in ('inserted','updated','unchanged','removed','duplicates')},
                             {'inserted': 1, 'updated': 1, 'unchanged': 1, 'removed': 1, 'duplicates': 0})
            with self.assertRaises(PermissionError):
                prepare_saved_sync(db, 2, sync_id, 'full')
            result = apply_saved_sync(db, 1, sync_id, Path(self.tmp.name) / 'backups')
            self.assertTrue(Path(result['backup']).is_file())
            rows = {r['phone_normalized']: r for r in db.execute('SELECT * FROM inventory')}
            self.assertEqual((rows[PHONE]['sale_status'], rows[PHONE]['phone_display'], rows[PHONE]['active']),
                             ('sold', '0900 000 0000', 1))
            self.assertEqual((rows[OTHER]['sale_status'], rows[OTHER]['active']), ('reserved', 1))
            self.assertEqual(rows[third]['active'], 0)
            self.assertEqual((rows[new]['sale_status'], rows[new]['active']), ('available', 1))
            batch = db.execute('SELECT mode,removed FROM import_batches ORDER BY id DESC LIMIT 1').fetchone()
            self.assertEqual(tuple(batch), ('full', 1))
        finally:
            db.close()

    def test_full_sync_is_atomic_and_stale_preview_is_rejected(self):
        db_path = Path(self.tmp.name) / 'atomic.db'
        db = connect(db_path)
        bootstrap_owner(db, 1)
        try:
            self.book([[PHONE, 1, 2, 10, None, 'همراه اول'], [OTHER, 1, 2, 20, None, 'همراه اول']])
            import_workbook(db, 1, self.path)
            self.book([[PHONE, 1, 2, 10, None, 'همراه اول']])
            sync_id = save_sync_upload(db, 1, self.path)
            prepare_saved_sync(db, 1, sync_id, 'full')
            with patch('espadna.importer.audit', side_effect=RuntimeError('simulated sync failure')):
                with self.assertRaises(RuntimeError):
                    apply_saved_sync(db, 1, sync_id, Path(self.tmp.name) / 'backups')
            self.assertEqual(dict(db.execute('SELECT phone_normalized,active FROM inventory')), {PHONE: 1, OTHER: 1})
            self.assertEqual(db.execute("SELECT count(*) FROM import_batches WHERE mode='full'").fetchone()[0], 0)
            db.execute("UPDATE inventory SET operator='ایرانسل' WHERE phone_normalized=?", (PHONE,))
            with self.assertRaisesRegex(ValueError, 'changed after preview'):
                apply_saved_sync(db, 1, sync_id, Path(self.tmp.name) / 'backups')
        finally:
            db.close()

    def test_sync_diagnostics_include_sheet_and_row_and_duplicate_preview(self):
        row = [PHONE, 1, 2, 10, None, 'همراه اول']
        self.book([row], [row])
        sync_id = save_sync_upload(self.db, 1, self.path)
        report = prepare_saved_sync(self.db, 1, sync_id, 'full')
        self.assertEqual(report['duplicates'], 1)
        self.assertEqual(report['duplicate_rows'], [{'sheet': 'second', 'row': 2}])
        bad = [OTHER, 1, 2, 10, None, 'invalid operator']
        self.book([row], [bad])
        with self.assertRaisesRegex(ValueError, "sheet 'second', row 2"):
            save_sync_upload(self.db, 1, self.path)

    def test_full_sync_deactivation_invalidates_scheduled_content(self):
        from espadna.content import approve_plan, create_plan, inspect_plan
        db_path = Path(self.tmp.name) / 'content-sync.db'
        db = connect(db_path)
        bootstrap_owner(db, 1)
        try:
            phones = [f'090000000{i:02}' for i in range(6)]
            rows = [[phone, 1, 2, None, None, 'همراه اول'] for phone in phones]
            self.book(rows)
            import_workbook(db, 1, self.path)
            plan = create_plan(db, 1, '2026-11')
            approve_plan(db, 1, plan)
            self.assertTrue(inspect_plan(db, 1, plan)['ready'])
            self.book(rows[1:])
            sync_id = save_sync_upload(db, 1, self.path)
            prepare_saved_sync(db, 1, sync_id, 'full')
            apply_saved_sync(db, 1, sync_id, Path(self.tmp.name) / 'backups')
            status = inspect_plan(db, 1, plan)
            self.assertFalse(status['ready'])
            self.assertGreater(status['stale_slots'], 0)
            self.assertEqual(db.execute('SELECT active FROM inventory WHERE phone_normalized=?', (phones[0],)).fetchone()[0], 0)
        finally:
            db.close()

    def test_validation_atomicity(self):
        for index, value in [(0, 9000000000), (1, True), (1, 3), (2, 4), (3, -1), (3, '=1+1'), (5, 'unverified')]:
            with self.subTest(index=index, value=value):
                good = [PHONE, 1, 2, 0, None, 'همراه اول']
                bad = [OTHER, 1, 2, 0, None, 'همراه اول']
                bad[index] = value
                self.book([good, bad])
                with self.assertRaises(ValueError):
                    import_workbook(self.db, 1, self.path)
                self.assertEqual(self.db.execute('SELECT count(*) FROM inventory').fetchone()[0], 0)
                self.assertEqual(self.db.execute('SELECT count(*) FROM import_batches').fetchone()[0], 0)

    def test_dry_run_and_audit(self):
        self.book()
        self.assertEqual(import_workbook(self.db, 1, self.path, True)['inserted'], 1)
        self.assertEqual(self.db.execute('SELECT count(*) FROM inventory').fetchone()[0], 0)
        import_workbook(self.db, 1, self.path)
        self.assertEqual(self.db.execute("SELECT count(*) FROM audit WHERE action='import'").fetchone()[0], 1)
        with self.assertRaises(Exception):
            self.db.execute('DELETE FROM audit')

    def test_roles_and_last_owner(self):
        set_role(self.db, 1, 2, 'viewer')
        self.book()
        for actor in [2, 99]:
            with self.assertRaises(PermissionError):
                import_workbook(self.db, actor, self.path)
            with self.assertRaises(PermissionError):
                set_role(self.db, actor, 3, 'admin')
        with self.assertRaises(ValueError):
            set_role(self.db, 1, 1, 'viewer')
        set_role(self.db, 1, 3, 'admin')
        import_workbook(self.db, 3, self.path)
        set_sale_status(self.db, 3, PHONE, 'reserved')
        with self.assertRaises(PermissionError):
            set_sale_status(self.db, 2, PHONE, 'sold')

    def test_revocation_and_reactivation_preserve_history(self):
        from espadna.core import revoke_user, require
        set_role(self.db, 1, 2, 'admin')
        with self.assertRaises(PermissionError):
            revoke_user(self.db, 2, 1)
        with self.assertRaises(ValueError):
            revoke_user(self.db, 1, 1)
        revoke_user(self.db, 1, 2)
        with self.assertRaises(PermissionError):
            require(self.db, 2, {'admin'})
        set_role(self.db, 1, 2, 'viewer')
        self.assertEqual(require(self.db, 2, {'viewer'}), 'viewer')
        self.assertEqual(self.db.execute("SELECT count(*) FROM audit WHERE action='user_revoked'").fetchone()[0], 1)

    def test_public_boundary(self):
        self.book()
        import_workbook(self.db, 1, self.path)
        row = self.db.execute('SELECT * FROM inventory').fetchone()
        text = caption(row)
        self.assertNotIn('12345', text)
        self.assertNotIn('price', text)
        self.assertIn(PHONE, text)
        set_sale_status(self.db, 1, PHONE, 'reserved')
        with self.assertRaises(ValueError):
            caption(self.db.execute('SELECT * FROM inventory').fetchone())

    def test_phone_digits_and_spacing(self):
        self.assertEqual(phone_text('۰۹۰۰ ۰۰۰ ۰۰۰۰'), PHONE)
        self.assertEqual(phone_values('091000 90 380'), ('09100090380', '091000 90 380'))
        for value in ['9000000000', '09-000000000', '=09000000000']:
            with self.assertRaises(ValueError):
                phone_text(value)

    def test_all_614_display_values_round_trip_exactly(self):
        rows = []
        expected = {}
        for index in range(614):
            normalized = f'091{index:08d}'
            display = f'{normalized[:6]} {normalized[6:8]} {normalized[8:]}'
            rows.append([display, 1, 2, None, None, 'همراه اول'])
            expected[normalized] = display
        self.book(rows)
        result = import_workbook(self.db, 1, self.path)
        self.assertEqual((result['rows'], result['unique'], result['inserted']), (614, 614, 614))
        stored = dict(self.db.execute('SELECT phone_normalized,phone_display FROM inventory'))
        self.assertEqual(stored, expected)

    def test_database_error_rolls_back_inventory_batch_and_audit(self):
        self.book()
        before = self.db.execute('SELECT count(*) FROM audit').fetchone()[0]
        with patch('espadna.importer.audit', side_effect=RuntimeError('simulated write failure')):
            with self.assertRaises(RuntimeError):
                import_workbook(self.db, 1, self.path)
        for table in ('inventory', 'import_batches'):
            self.assertEqual(self.db.execute('SELECT count(*) FROM ' + table).fetchone()[0], 0)
        self.assertEqual(self.db.execute('SELECT count(*) FROM audit').fetchone()[0], before)

    def test_blank_operator_is_not_inferred(self):
        self.book([[PHONE, 1, 2, None, None, None]])
        import_workbook(self.db, 1, self.path)
        from espadna.core import public_item
        result = public_item(self.db.execute('SELECT * FROM inventory').fetchone())
        self.assertNotIn('operator', result)
        self.assertEqual(set(result), {'phone', 'sim_type', 'sim_status'})

    def test_internal_price_is_exact_and_bounded(self):
        from espadna.importer import price
        self.assertEqual(price('123456789012345678901234.123456'), '123456789012345678901234.123456')
        self.assertEqual(price('123.4500'), '123.45')
        self.assertEqual(price('0.000000'), '0')
        for value in ('1e999999999', '1e-999999999', 'NaN', 'Infinity'):
            with self.subTest(value=value), self.assertRaises(ValueError):
                price(value)

    def test_sales_terms_only_row_is_completely_ignored(self):
        self.book([[PHONE, 1, 2, 0, None, None], [None, None, None, None, '=1+1', None]])
        result = import_workbook(self.db, 1, self.path)
        self.assertEqual((result['rows'], result['inserted']), (1, 1))

    def test_archive_and_row_limits_and_missing_headers(self):
        self.book()
        with patch('espadna.importer.MAX_BYTES', 1):
            with self.assertRaises(ValueError):
                parse_workbook(self.path)
        with patch('espadna.importer.MAX_UNPACKED', 1):
            with self.assertRaises(ValueError):
                parse_workbook(self.path)
        with patch('espadna.importer.MAX_ROWS', 0):
            with self.assertRaises(ValueError):
                parse_workbook(self.path)
        wb = openpyxl.Workbook()
        wb.active.append(['unknown'])
        wb.save(self.path)
        wb.close()
        with self.assertRaises(ValueError):
            import_workbook(self.db, 1, self.path)

    def test_bot_denies_before_download(self):
        from espadna.bot import build_application
        from types import SimpleNamespace
        import asyncio
        dbpath = Path(self.tmp.name) / 'bot.db'
        db = connect(dbpath)
        bootstrap_owner(db, 1)
        set_role(db, 1, 2, 'viewer')
        db.close()
        app = build_application('123456:TEST_ONLY_PLACEHOLDER', dbpath)
        callback = app.handlers[0][1].callback
        for actor, chat_type in [(2, 'private'), (99, 'private'), (1, 'group')]:
            msg = SimpleNamespace(reply_text=AsyncMock(), document=None)
            update = SimpleNamespace(effective_chat=SimpleNamespace(type=chat_type), effective_user=SimpleNamespace(id=actor), effective_message=msg)
            context = SimpleNamespace(bot=SimpleNamespace(get_file=AsyncMock()))
            asyncio.run(callback(update, context))
            context.bot.get_file.assert_not_called()
            msg.reply_text.assert_awaited_once()


if __name__ == '__main__':
    unittest.main()
